import io

from pptx import Presentation
from pptx.util import Inches, Pt

from src.reports.pdf_report import DISCLAIMER
from src.reports.report_context import ReportContext

_TITLE_LAYOUT = 0
_TITLE_CONTENT_LAYOUT = 1
_TITLE_ONLY_LAYOUT = 5


def _add_bullet_slide(prs: Presentation, title: str, bullets: list[str]):
    slide = prs.slides.add_slide(prs.slide_layouts[_TITLE_CONTENT_LAYOUT])
    slide.shapes.title.text = title

    body = slide.placeholders[1].text_frame
    body.clear()
    for i, bullet in enumerate(bullets or ["(nothing to report)"]):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = bullet
        p.font.size = Pt(16)
    return slide


def _add_table_slide(prs: Presentation, title: str, header: list[str], rows: list[list[str]]):
    slide = prs.slides.add_slide(prs.slide_layouts[_TITLE_ONLY_LAYOUT])
    slide.shapes.title.text = title

    if not rows:
        return slide

    n_rows, n_cols = len(rows) + 1, len(header)
    table_shape = slide.shapes.add_table(n_rows, n_cols, Inches(0.4), Inches(1.4), Inches(9), Inches(0.4 * n_rows))
    table = table_shape.table

    for c, label in enumerate(header):
        table.cell(0, c).text = label
        for run in table.cell(0, c).text_frame.paragraphs[0].runs:
            run.font.bold = True
            run.font.size = Pt(12)

    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row):
            table.cell(r, c).text = str(value)
            for run in table.cell(r, c).text_frame.paragraphs[0].runs:
                run.font.size = Pt(10)

    return slide


def generate_pptx_report(ctx: ReportContext) -> bytes:
    prs = Presentation()

    _title_slide(prs, ctx)

    if ctx.synthesis is not None:
        _add_bullet_slide(prs, "Executive Summary", [ctx.synthesis.overall_assessment] + ctx.synthesis.key_findings)

        bp = ctx.synthesis.business_performance
        _add_bullet_slide(
            prs,
            "Business Performance",
            [f"Revenue: {bp.revenue}", f"Growth: {bp.growth}", f"Margin: {bp.margin}", f"Cash: {bp.cash}"]
            + bp.key_operational_metrics,
        )
    else:
        _add_bullet_slide(
            prs, "Executive Summary",
            ["No Deep Analysis run has completed for this project yet — showing available data only."],
        )

    if ctx.concerns:
        _add_bullet_slide(
            prs, "Key Concerns",
            [f"{c.title} ({c.severity}): {c.business_impact or 'impact not stated'}" for c in ctx.concerns],
        )

    if ctx.opportunities:
        _add_bullet_slide(prs, "Key Opportunities", [f"{o.title}: {o.rationale}" for o in ctx.opportunities])

    if ctx.synthesis is not None and ctx.synthesis.strategic_options:
        _add_bullet_slide(
            prs, "Strategic Options",
            [f"{o.option} ({o.feasibility.value} feasibility): {o.recommendation}" for o in ctx.synthesis.strategic_options],
        )

    if ctx.synthesis is not None and ctx.synthesis.ninety_day_plan:
        _add_table_slide(
            prs,
            "90-Day Action Plan",
            ["Action", "KPI", "Expected impact"],
            [[item.action, item.kpi or "—", item.expected_impact or "—"] for item in ctx.synthesis.ninety_day_plan],
        )

    if ctx.synthesis is not None and ctx.synthesis.missing_information:
        _add_bullet_slide(prs, "Missing Information", ctx.synthesis.missing_information)

    return _to_bytes(prs)


def _title_slide(prs: Presentation, ctx: ReportContext):
    slide = prs.slides.add_slide(prs.slide_layouts[_TITLE_LAYOUT])
    slide.shapes.title.text = ctx.company.name
    subtitle = slide.placeholders[1]
    subtitle.text = f"{ctx.project.name}\nBusiness Consulting Report — {ctx.generated_at.strftime('%B %d, %Y')}\n\n{DISCLAIMER}"
    for paragraph in subtitle.text_frame.paragraphs:
        paragraph.font.size = Pt(14)
    return slide


def _to_bytes(prs: Presentation) -> bytes:
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
