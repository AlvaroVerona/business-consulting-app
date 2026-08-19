import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from src.reports.report_context import ReportContext

_HEADER_FONT = Font(bold=True)


def _write_table(ws: Worksheet, header: list[str], rows: list[list]) -> None:
    ws.append(header)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append(row)
    for column_cells in ws.columns:
        length = max((len(str(c.value)) for c in column_cells if c.value is not None), default=10)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 60)


def generate_excel_export(ctx: ReportContext) -> bytes:
    wb = Workbook()

    financial_ws = wb.active
    financial_ws.title = "Financial Analysis"
    _write_table(
        financial_ws,
        ["Period", "Revenue", "COGS", "Opex", "EBITDA", "EBITDA implied?", "Gross Margin", "EBITDA Margin", "Opex Ratio"],
        [
            [
                p.period, p.revenue, p.cogs, p.opex, p.ebitda, "Yes" if p.ebitda_is_implied else "No",
                p.gross_margin, p.ebitda_margin, p.opex_ratio,
            ]
            for p in ctx.financial_analysis.periods
        ],
    )

    concerns_ws = wb.create_sheet("Concerns")
    _write_table(
        concerns_ws,
        ["ID", "Title", "Severity", "Confidence", "Business Impact", "Recommended Action"],
        [[c.id, c.title, c.severity, c.confidence, c.business_impact, c.recommended_action] for c in ctx.concerns],
    )

    opportunities_ws = wb.create_sheet("Opportunities")
    _write_table(
        opportunities_ws,
        ["ID", "Title", "Confidence", "Rationale", "Next Step", "Risks"],
        [[o.id, o.title, o.confidence, o.rationale, o.next_step, o.risks] for o in ctx.opportunities],
    )

    hypotheses_ws = wb.create_sheet("Hypotheses")
    _write_table(
        hypotheses_ws,
        ["ID", "Statement", "Status", "Origin", "Priority", "Data Required", "Next Test"],
        [
            [h.id, h.statement, h.status, h.origin, h.priority, h.data_required, h.next_test]
            for h in ctx.hypotheses
        ],
    )

    return _to_bytes(wb)


def _to_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
