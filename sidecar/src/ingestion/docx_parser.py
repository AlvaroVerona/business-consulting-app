from docx import Document as DocxDocument


class DocxParser:
    """One chunk per non-empty paragraph. Tables are flattened row-by-row
    (python-docx exposes no page numbers — DOCX has none until rendered)."""

    def parse(self, path: str) -> list[dict]:
        doc = DocxDocument(path)
        chunks = []

        for i, paragraph in enumerate(doc.paragraphs):
            if paragraph.text.strip():
                chunks.append(
                    {"content": paragraph.text.strip(), "location": {"paragraph": i + 1}}
                )

        for t, table in enumerate(doc.tables):
            for r, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    chunks.append(
                        {
                            "content": " | ".join(cells),
                            "location": {"table": t + 1, "row": r + 1},
                        }
                    )

        return chunks
