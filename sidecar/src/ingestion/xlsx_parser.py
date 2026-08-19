import openpyxl


class XLSXParser:
    """One chunk per row per sheet, with computed values (data_only=True)
    rather than formula text — spec section 6: 'perform real calculations
    rather than relying only on text extraction'. The formula itself is
    kept in `location` so a Finding can still show its derivation."""

    def parse(self, path: str) -> list[dict]:
        values_wb = openpyxl.load_workbook(path, data_only=True)
        formulas_wb = openpyxl.load_workbook(path, data_only=False)
        chunks = []

        for sheet_name in values_wb.sheetnames:
            values_ws = values_wb[sheet_name]
            formulas_ws = formulas_wb[sheet_name]

            header = [cell.value for cell in values_ws[1]] if values_ws.max_row >= 1 else []

            for row_idx, row in enumerate(values_ws.iter_rows(min_row=2), start=2):
                cells = []
                formulas = {}

                for col_idx, cell in enumerate(row):
                    formula_cell = formulas_ws.cell(row=row_idx, column=col_idx + 1)
                    is_formula = isinstance(formula_cell.value, str) and formula_cell.value.startswith("=")

                    if cell.value is None and not is_formula:
                        continue

                    col_name = header[col_idx] if col_idx < len(header) and header[col_idx] else cell.coordinate
                    # data_only=True has no cached value until the file has been opened and
                    # saved by Excel/LibreOffice; fall back to the formula text so nothing is lost.
                    display_value = cell.value if cell.value is not None else formula_cell.value
                    cells.append(f"{col_name}: {display_value}")

                    if is_formula:
                        formulas[cell.coordinate] = formula_cell.value

                if not cells:
                    continue

                chunks.append(
                    {
                        "content": "; ".join(cells),
                        "location": {
                            "sheet": sheet_name,
                            "row": row_idx,
                            "formulas": formulas or None,
                        },
                    }
                )

        return chunks
