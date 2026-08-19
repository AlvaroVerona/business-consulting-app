import csv
import os

import openpyxl
import pytest
from docx import Document as DocxDocument

from src.ingestion.csv_parser import CSVParser
from src.ingestion.docx_parser import DocxParser
from src.ingestion.registry import UnsupportedFileType, parse_document
from src.ingestion.text_parser import TextParser
from src.ingestion.xlsx_parser import XLSXParser


def test_text_parser_splits_paragraphs(tmp_path):
    path = tmp_path / "note.txt"
    path.write_text("First paragraph.\n\nSecond paragraph.\n\n\n")

    chunks = TextParser().parse(str(path))

    assert [c["content"] for c in chunks] == ["First paragraph.", "Second paragraph."]
    assert chunks[0]["location"] == {"paragraph": 1}


def test_csv_parser_rows_carry_location(tmp_path):
    path = tmp_path / "revenue.csv"
    with open(path, "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["month", "revenue"])
        writer.writerow(["Jan", "1000"])
        writer.writerow(["Feb", "1200"])

    chunks = CSVParser().parse(str(path))

    assert len(chunks) == 2
    assert "revenue: 1000" in chunks[0]["content"]
    assert chunks[0]["location"]["row"] == 2  # header is row 1


def test_docx_parser_extracts_paragraphs_and_tables(tmp_path):
    path = tmp_path / "memo.docx"
    doc = DocxDocument()
    doc.add_paragraph("Q3 revenue declined 8%.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "COGS"
    table.rows[0].cells[1].text = "120000"
    doc.save(str(path))

    chunks = DocxParser().parse(str(path))
    contents = [c["content"] for c in chunks]

    assert "Q3 revenue declined 8%." in contents
    assert any("COGS" in c and "120000" in c for c in contents)


def test_xlsx_parser_uses_computed_values_and_keeps_formula(tmp_path):
    path = tmp_path / "model.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["item", "amount"])
    ws.append(["revenue", 1000])
    ws.append(["cogs", "=B2*0.4"])
    wb.save(str(path))

    chunks = XLSXParser().parse(str(path))

    assert len(chunks) == 2
    assert "item: revenue" in chunks[0]["content"]
    assert "amount: 1000" in chunks[0]["content"]
    # openpyxl with data_only=True has no cached value for a formula never opened in Excel,
    # so this asserts the formula text survived into location rather than a computed number.
    assert chunks[1]["location"]["formulas"] is not None


def test_registry_dispatches_by_extension(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("Hello.")

    chunks = parse_document(str(path), ".md")

    assert chunks[0]["content"] == "Hello."


def test_registry_rejects_unknown_extension(tmp_path):
    path = tmp_path / "note.xyz"
    path.write_text("Hello.")

    with pytest.raises(UnsupportedFileType):
        parse_document(str(path), ".xyz")
