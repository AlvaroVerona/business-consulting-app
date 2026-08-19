import io

from pptx import Presentation

from src.database.repository import Repository
from src.reports.pptx_report import generate_pptx_report
from src.reports.report_context import build_report_context
from tests.report_fixtures import seed_full_project


def test_full_deck_has_expected_slides(db_session):
    repo = Repository(db_session)
    project_id = seed_full_project(repo)
    ctx = build_report_context(repo, project_id)

    pptx_bytes = generate_pptx_report(ctx)

    prs = Presentation(io.BytesIO(pptx_bytes))
    titles = [s.shapes.title.text for s in prs.slides if s.shapes.title is not None]

    assert titles[0] == "Acme Wine Bar"
    assert "Executive Summary" in titles
    assert "Key Concerns" in titles
    assert "Key Opportunities" in titles
    assert "90-Day Action Plan" in titles


def test_deck_with_no_completed_run_does_not_crash(db_session):
    repo = Repository(db_session)
    company = repo.create_company(name="Acme")
    project = repo.create_project(company.id, name="P1")
    ctx = build_report_context(repo, project.id)

    pptx_bytes = generate_pptx_report(ctx)

    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) >= 1
