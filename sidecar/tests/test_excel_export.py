import io

from openpyxl import load_workbook

from src.database.repository import Repository
from src.reports.excel_export import generate_excel_export
from src.reports.report_context import build_report_context
from tests.report_fixtures import seed_full_project


def test_export_produces_expected_sheets_and_rows(db_session):
    repo = Repository(db_session)
    project_id = seed_full_project(repo)
    ctx = build_report_context(repo, project_id)

    xlsx_bytes = generate_excel_export(ctx)

    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert wb.sheetnames == ["Financial Analysis", "Concerns", "Opportunities", "Hypotheses"]

    financial_ws = wb["Financial Analysis"]
    assert financial_ws["A1"].value == "Period"
    assert financial_ws.max_row == 4  # header + 3 periods

    concerns_ws = wb["Concerns"]
    assert concerns_ws.cell(row=2, column=2).value == "Gross margin compression"


def test_export_with_no_data_still_produces_valid_workbook(db_session):
    repo = Repository(db_session)
    company = repo.create_company(name="Acme")
    project = repo.create_project(company.id, name="P1")
    ctx = build_report_context(repo, project.id)

    xlsx_bytes = generate_excel_export(ctx)

    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert wb["Financial Analysis"].max_row == 1  # header only
